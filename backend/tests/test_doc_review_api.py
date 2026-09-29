"""AI 文档审阅（`doc-review` 能力）的解析 / 审阅 / 对话 / 导出链路测试。

关注四件事：
1. 解析：docx → 段落块按文档顺序（含表格内段落、去重合并单元格），空段不入块；
2. 审阅与对话：模型输出不可信 —— 未知块 id、空文本、超长膨胀都要被收敛掉；
3. 模型归属与 slides 同一套硬约束（只用自己的模型）；
4. 导出：blocks 回写**原始 docx**，改过的段落文本替换、未改段落不动，样式保留。
"""

from __future__ import annotations

from io import BytesIO
from types import SimpleNamespace
from typing import Any

import pytest
from docx import Document
from fastapi import HTTPException, UploadFile
from sqlalchemy.pool import StaticPool
from sqlmodel import Session, SQLModel, create_engine

from app.api import agent_apps
from app.api.agent_apps import (
    DocChatRequest,
    DocExportRequest,
    DocReviewRequest,
    chat_doc,
    export_doc,
    parse_doc,
    review_doc,
)
from app.core.doc_review import DocBlock, sanitize_actions, sanitize_issues, store_document
from app.db.models import ModelConfig, Tenant, User


def _test_session() -> Session:
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    SQLModel.metadata.create_all(engine)
    return Session(engine)


def _member(user_id: str = "user_member") -> User:
    return User(
        id=user_id,
        tenant_id="tenant_demo",
        username=user_id,
        role="member",
        password_hash="test",
    )


def _prepare(db: Session) -> None:
    db.add(Tenant(id="tenant_demo", name="Demo"))
    db.add(_member())
    db.add(
        ModelConfig(
            id="model-1",
            tenant_id="tenant_demo",
            user_id="user_member",
            name="model-1",
            model="deepseek-v4-flash",
            api_key_encrypted="encrypted",
            enabled=True,
        )
    )
    db.commit()


def _patch_runtime_model(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(
        agent_apps,
        "resolve_model_config_for_runtime",
        lambda db, tenant_id, config_id: SimpleNamespace(id=config_id, tenant_id=tenant_id, model="deepseek-v4-flash"),
    )


class _FakeLLMClient:
    calls: list[dict[str, Any]] = []
    response: Any = {}

    def __init__(self, model_config: Any) -> None:
        self.model_config = model_config

    def generate_json(self, system_prompt: str, user_payload: dict[str, Any], **_: Any) -> Any:
        type(self).calls.append({"system": system_prompt, "payload": user_payload})
        return type(self).response


@pytest.fixture(autouse=True)
def _reset_fake() -> None:
    _FakeLLMClient.calls = []
    _FakeLLMClient.response = {}


# ---------------------------------------------------------------- 样例文档


def _build_docx() -> bytes:
    document = Document()
    document.add_heading("服务协议", level=1)
    document.add_paragraph("甲方：上海某某科技有限公司。")
    document.add_paragraph("乙方负责违约责任的约定存在错别字：违约为本。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "条款"
    table.cell(0, 1).text = "内容"
    table.cell(1, 0).text = "保密期"
    table.cell(1, 1).text = "自终止日起 5 个工作日"
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _upload(data: bytes, filename: str = "服务协议.docx") -> UploadFile:
    return UploadFile(file=BytesIO(data), filename=filename)


def _blocks_from(payload: dict[str, Any]) -> list[dict[str, Any]]:
    return payload["blocks"]


# ---------------------------------------------------------------- 解析


def test_parse_extracts_blocks_in_document_order_including_table() -> None:
    with _test_session() as db:
        _prepare(db)
        payload = parse_doc("tenant_demo", _upload(_build_docx()), db, _member())

    blocks = _blocks_from(payload)
    texts = [block["text"] for block in blocks]
    # 标题、正文、表格内容都在，且保持文档顺序
    assert texts[0] == "服务协议"
    assert "违约为本" in " ".join(texts)
    assert "保密期" in texts
    assert "自终止日起 5 个工作日" in texts
    # 表格块带 in_table 标记
    assert any(block["in_table"] for block in blocks)
    assert not blocks[0]["in_table"]
    assert blocks[0]["kind"] == "heading"
    # 块 id 连续稳定
    assert [block["id"] for block in blocks] == [f"p{i}" for i in range(len(blocks))]


def test_parse_rejects_non_docx_extension() -> None:
    with _test_session() as db:
        _prepare(db)
        with pytest.raises(HTTPException) as excinfo:
            parse_doc("tenant_demo", _upload(b"not word", "legacy.doc"), db, _member())
    assert ".docx" in str(excinfo.value.detail)


def test_parse_rejects_corrupted_file() -> None:
    with _test_session() as db:
        _prepare(db)
        with pytest.raises(HTTPException) as excinfo:
            parse_doc("tenant_demo", _upload(b"PK fake zip", "broken.docx"), db, _member())
    assert "无法解析" in str(excinfo.value.detail)


# ---------------------------------------------------------------- 审阅


def test_review_returns_sanitized_issues(monkeypatch: pytest.MonkeyPatch) -> None:
    _patch_runtime_model(monkeypatch)
    monkeypatch.setattr(agent_apps, "LLMClient", _FakeLLMClient)
    with _test_session() as db:
        _prepare(db)
        payload = parse_doc("tenant_demo", _upload(_build_docx()), db, _member())
        blocks = [DocBlock(**row) for row in _blocks_from(payload)]
        _FakeLLMClient.response = {
            "issues": [
                {
                    "block_ids": ["p2", "p999"],  # p999 不存在，应被丢掉
                    "type": "错别字",
                    "title": "「违约为本」应为「违约为准」",
                    "detail": "固定搭配是「以…为准」。",
                    "fixes": [
                        {"block_id": "p2", "new_text": "乙方负责违约责任的约定以本协议为准。"},
                        {"block_id": "p888", "new_text": "ghost"},  # 未知块，应被丢掉
                        {"block_id": "p2", "new_text": ""},  # 空文本，应被丢掉
                    ],
                },
                {"title": ""},  # 无标题，整条丢弃
                "not-a-dict",
            ]
        }
        result = review_doc(
            DocReviewRequest(tenant_id="tenant_demo", model_config_id="model-1", doc_id=payload["doc_id"], blocks=blocks),
            db,
            _member(),
        )

    assert len(result["issues"]) == 1
    issue = result["issues"][0]
    assert issue["block_ids"] == ["p2"]
    assert len(issue["fixes"]) == 1
    assert issue["fixes"][0]["block_id"] == "p2"
    assert _FakeLLMClient.calls, "审阅必须调过模型"


def test_review_rejects_model_not_owned(monkeypatch: pytest.MonkeyPatch) -> None:
    _patch_runtime_model(monkeypatch)
    with _test_session() as db:
        _prepare(db)
        payload = parse_doc("tenant_demo", _upload(_build_docx()), db, _member())
        blocks = [DocBlock(**row) for row in _blocks_from(payload)]
        with pytest.raises(HTTPException) as excinfo:
            review_doc(
                DocReviewRequest(
                    tenant_id="tenant_demo", model_config_id="model-someone-else", doc_id=payload["doc_id"], blocks=blocks
                ),
                db,
                _member(),
            )
    assert excinfo.value.status_code == 403


# ---------------------------------------------------------------- 对话


def test_chat_drops_invalid_actions(monkeypatch: pytest.MonkeyPatch) -> None:
    _patch_runtime_model(monkeypatch)
    monkeypatch.setattr(agent_apps, "LLMClient", _FakeLLMClient)
    with _test_session() as db:
        _prepare(db)
        payload = parse_doc("tenant_demo", _upload(_build_docx()), db, _member())
        blocks = [DocBlock(**row) for row in _blocks_from(payload)]
        _FakeLLMClient.response = {
            "reply": "已修正错别字并归纳了保密条款。",
            "actions": [
                {"block_id": "p2", "new_text": "乙方负责违约责任的约定以本协议为准。", "reason": "修正错别字"},
                {"block_id": "p404", "new_text": "ghost"},
                {"block_id": "p3", "new_text": ""},
            ],
        }
        result = chat_doc(
            DocChatRequest(
                tenant_id="tenant_demo",
                model_config_id="model-1",
                doc_id=payload["doc_id"],
                message="把错别字改掉",
                history=[],
                blocks=blocks,
            ),
            db,
            _member(),
        )

    assert "错别字" in result["reply"]
    assert [action["block_id"] for action in result["actions"]] == ["p2"]


def test_sanitize_actions_rejects_runaway_inflation() -> None:
    blocks = [DocBlock(id="p0", text="短句。")]
    _FakeLLMClient.response = {}
    outcome = sanitize_actions(
        {"reply": "r", "actions": [{"block_id": "p0", "new_text": "长" * 5000}]}, blocks
    )
    assert outcome.actions == []


def test_sanitize_issues_ignores_malformed_payload() -> None:
    blocks = [DocBlock(id="p0", text="正文")]
    assert sanitize_issues(None, blocks) == []
    assert sanitize_issues({}, blocks) == []
    assert sanitize_issues({"issues": "not-a-list"}, blocks) == []


# ---------------------------------------------------------------- 导出


def test_export_writes_edits_back_into_original_docx() -> None:
    with _test_session() as db:
        _prepare(db)
        payload = parse_doc("tenant_demo", _upload(_build_docx()), db, _member())
        blocks = [DocBlock(**row) for row in _blocks_from(payload)]
        doc_id = payload["doc_id"]

        # 模拟用户应用了一条修改：p2 换成修正后的文本，其余不动
        edited = [
            DocBlock(**{**block.model_dump(), "text": "乙方负责违约责任的约定以本协议为准。"})
            if block.id == "p2"
            else block
            for block in blocks
        ]
        response = export_doc(
            DocExportRequest(tenant_id="tenant_demo", doc_id=doc_id, blocks=edited, file_name="服务协议"),
            db,
            _member(),
        )

    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert response.headers["x-doc-changed"] == "1"

    document = Document(BytesIO(response.body))
    texts = [paragraph.text for paragraph in document.paragraphs]
    assert "乙方负责违约责任的约定以本协议为准。" in texts
    assert not any("违约为本" in text for text in texts)
    assert "甲方：上海某某科技有限公司。" in texts  # 未修改段落保持原样

    # 标题样式保留（段落级回写不动样式）
    headings = [p for p in document.paragraphs if (p.style.name or "").startswith("Heading")]
    assert any(p.text == "服务协议" for p in headings)

    # 表格内容也回写
    table = document.tables[0]
    assert table.cell(1, 1).text == "自终止日起 5 个工作日"


def test_export_requires_stored_document() -> None:
    with _test_session() as db:
        _prepare(db)
        with pytest.raises(HTTPException) as excinfo:
            export_doc(
                DocExportRequest(
                    tenant_id="tenant_demo",
                    doc_id="missing",
                    blocks=[DocBlock(id="p0", text="x")],
                ),
                db,
                _member(),
            )
    assert excinfo.value.status_code == 404


def _ai_comments(document: Document) -> list[tuple[str, str]]:
    """(作者, 文本) 形式收集文档里的批注。"""
    return [(comment.author, comment.text) for comment in document.comments]


def test_export_adds_track_changes_comments_without_touching_text() -> None:
    with _test_session() as db:
        _prepare(db)
        payload = parse_doc("tenant_demo", _upload(_build_docx()), db, _member())
        blocks = [DocBlock(**row) for row in _blocks_from(payload)]
        commented = [
            DocBlock(**{**block.model_dump(), "comment": "【错别字】「违约为本」应为「违约为准」"})
            if block.id == "p2"
            else block
            for block in blocks
        ]
        response = export_doc(
            DocExportRequest(tenant_id="tenant_demo", doc_id=payload["doc_id"], blocks=commented),
            db,
            _member(),
        )

    document = Document(BytesIO(response.body))
    # 原文一个字都没动（批注模式不改文字）
    texts = [paragraph.text for paragraph in document.paragraphs]
    assert any("违约为本" in text for text in texts)
    # 批注存在，作者与内容正确
    comments = _ai_comments(document)
    assert ("AI 审阅", "【错别字】「违约为本」应为「违约为准」") in comments
    # 批注锚定在 p2 段（该段包含批注引用）
    from docx.oxml.ns import qn

    anchored = []
    for paragraph in document.paragraphs:
        if paragraph._p.find(".//" + qn("w:commentReference")) is not None:
            anchored.append(paragraph.text)
    assert any("违约为本" in text for text in anchored)


def test_export_comments_are_idempotent_and_clearable() -> None:
    with _test_session() as db:
        _prepare(db)
        payload = parse_doc("tenant_demo", _upload(_build_docx()), db, _member())
        blocks = [DocBlock(**row) for row in _blocks_from(payload)]
        doc_id = payload["doc_id"]

        def _with_comment(comment: str) -> list[DocBlock]:
            return [
                DocBlock(**{**block.model_dump(), "comment": comment})
                if block.id == "p2"
                else block
                for block in blocks
            ]

        # 同一批 blocks 连续导出两次：批注不得重复累积
        first = export_doc(
            DocExportRequest(tenant_id="tenant_demo", doc_id=doc_id, blocks=_with_comment("第一次批注")),
            db,
            _member(),
        )
        second = export_doc(
            DocExportRequest(tenant_id="tenant_demo", doc_id=doc_id, blocks=_with_comment("第二次批注")),
            db,
            _member(),
        )
        document = Document(BytesIO(second.body))
        assert _ai_comments(document).count(("AI 审阅", "第二次批注")) == 1
        assert ("AI 审阅", "第一次批注") not in _ai_comments(document)
        assert first.status_code == 200

        # 批注清空后再导出：AI 批注消失，用户自己的批注不受影响
        plain = export_doc(
            DocExportRequest(tenant_id="tenant_demo", doc_id=doc_id, blocks=blocks),
            db,
            _member(),
        )
        assert _ai_comments(Document(BytesIO(plain.body))) == []


def test_store_document_evicts_oldest_beyond_capacity(monkeypatch: pytest.MonkeyPatch) -> None:
    from app.core import doc_review

    monkeypatch.setattr(doc_review, "_STORE_MAX", 3)
    doc_review._DOC_STORE.clear()
    ids = [store_document(f"d{i}.docx", object(), [object()]) for i in range(5)]
    assert len(doc_review._DOC_STORE) == 3
    assert ids[-1] in doc_review._DOC_STORE
    assert ids[0] not in doc_review._DOC_STORE
    doc_review._DOC_STORE.clear()
