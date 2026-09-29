"""Agent 广场（`/api/enterprise/agent-apps`）的生成与导出链路测试。

关注四件事：
1. 模型归属：只能用当前用户自己配置的模型，缺失/越权/停用都要挡住，且**不回落**租户默认模型；
2. 选项对齐：左侧勾了骨架页就得有、没勾的不许出现、自定义页数要截断（「选了首页就得有个首页」）；
3. 模型输出不可信：版式越界、条目超量、骨架页漏生成都要被收敛；
4. 导出：deck 能落成可打开的 pptx，空 deck 被挡住。
"""

from __future__ import annotations

from io import BytesIO
from types import SimpleNamespace
from typing import Any

import pytest
from fastapi import HTTPException
from pptx import Presentation
from sqlalchemy.pool import StaticPool
from sqlmodel import Session, SQLModel, create_engine

from app.api import agent_apps
from app.api.agent_apps import (
    PPTX_MEDIA_TYPE,
    SlidesExportRequest,
    SlidesGenerateRequest,
    export_slides,
    generate_slides,
)
from app.core import slides_deck
from app.db.models import ModelConfig, Tenant, User


def _test_session() -> Session:
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    SQLModel.metadata.create_all(engine)
    return Session(engine)


def _member(user_id: str = "user_member") -> User:
    return User(
        id=user_id,
        tenant_id="tenant_demo",
        username=user_id,
        role="member",
        password_hash="test",
    )


def _model_config(model_id: str, *, user_id: str | None, enabled: bool = True) -> ModelConfig:
    return ModelConfig(
        id=model_id,
        tenant_id="tenant_demo",
        user_id=user_id,
        name=f"model-{model_id}",
        model="deepseek-v4-flash",
        api_key_encrypted="encrypted",
        enabled=enabled,
    )


class _FakeLLMClient:
    """记录入参并回放固定 JSON 的假客户端。"""

    calls: list[dict[str, Any]] = []
    response: Any = {"deck_title": "库存收尾", "page_label": "库存收尾", "pages": []}

    def __init__(self, model_config: Any) -> None:
        self.model_config = model_config

    def generate_json(self, system_prompt: str, user_payload: dict[str, Any], **_: Any) -> Any:
        type(self).calls.append({"system": system_prompt, "payload": user_payload})
        return type(self).response


@pytest.fixture(autouse=True)
def _reset_fake() -> None:
    _FakeLLMClient.calls = []
    _FakeLLMClient.response = {"deck_title": "库存收尾", "page_label": "库存收尾", "pages": []}


def _prepare(db: Session, *, owner_id: str | None = "user_member", enabled: bool = True) -> None:
    db.add(Tenant(id="tenant_demo", name="Demo"))
    db.add(_member())
    db.add(_model_config("model-1", user_id=owner_id, enabled=enabled))
    db.commit()


def _request(**overrides: Any) -> SlidesGenerateRequest:
    payload: dict[str, Any] = {
        "tenant_id": "tenant_demo",
        "model_config_id": "model-1",
        "narrative": "本月完成库存收尾，两条主线：合同转签与开票回收。",
    }
    payload.update(overrides)
    return SlidesGenerateRequest(**payload)


def _patch_runtime_model(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(
        agent_apps,
        "resolve_model_config_for_runtime",
        lambda db, tenant_id, config_id: SimpleNamespace(id=config_id, tenant_id=tenant_id, model="deepseek-v4-flash"),
    )


def _patch_llm(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(agent_apps, "LLMClient", _FakeLLMClient)


# ---------------------------------------------------------------- 模型归属


def test_generate_rejects_model_owned_by_another_user(monkeypatch: pytest.MonkeyPatch) -> None:
    _patch_llm(monkeypatch)
    with _test_session() as db:
        _prepare(db, owner_id="user_someone_else")
        with pytest.raises(HTTPException) as excinfo:
            generate_slides(_request(), db, _member())
    assert excinfo.value.status_code == 403
    assert "你在「模型配置」里配置的模型" in str(excinfo.value.detail)
    assert _FakeLLMClient.calls == []


def test_generate_rejects_user_shared_tenant_model(monkeypatch: pytest.MonkeyPatch) -> None:
    """租户级模型（user_id 为空）不属于任何人，同样不允许被借用。"""
    _patch_llm(monkeypatch)
    with _test_session() as db:
        _prepare(db, owner_id=None)
        with pytest.raises(HTTPException) as excinfo:
            generate_slides(_request(), db, _member())
    assert excinfo.value.status_code == 403
    assert _FakeLLMClient.calls == []


def test_generate_rejects_disabled_model(monkeypatch: pytest.MonkeyPatch) -> None:
    _patch_llm(monkeypatch)
    with _test_session() as db:
        _prepare(db, enabled=False)
        with pytest.raises(HTTPException) as excinfo:
            generate_slides(_request(), db, _member())
    assert excinfo.value.status_code == 400
    assert "已停用" in str(excinfo.value.detail)


def test_generate_requires_model_config_id(monkeypatch: pytest.MonkeyPatch) -> None:
    _patch_llm(monkeypatch)
    with _test_session() as db:
        _prepare(db)
        with pytest.raises(HTTPException) as excinfo:
            generate_slides(_request(model_config_id=""), db, _member())
    assert excinfo.value.status_code == 400
    assert _FakeLLMClient.calls == []


# ---------------------------------------------------------------- 输入校验


def test_generate_rejects_blank_narrative(monkeypatch: pytest.MonkeyPatch) -> None:
    _patch_llm(monkeypatch)
    with _test_session() as db:
        _prepare(db)
        with pytest.raises(HTTPException) as excinfo:
            generate_slides(_request(narrative="   "), db, _member())
    assert excinfo.value.status_code == 400
    assert "正文口述" in str(excinfo.value.detail)
    assert _FakeLLMClient.calls == []


def test_generate_rejects_out_of_range_page_count(monkeypatch: pytest.MonkeyPatch) -> None:
    _patch_llm(monkeypatch)
    with _test_session() as db:
        _prepare(db)
        with pytest.raises(HTTPException) as excinfo:
            generate_slides(_request(page_count=0), db, _member())
    assert excinfo.value.status_code == 400
    assert _FakeLLMClient.calls == []


# ---------------------------------------------------------------- 正常链路


def test_generate_passes_user_model_and_returns_normalized_deck(monkeypatch: pytest.MonkeyPatch) -> None:
    _patch_runtime_model(monkeypatch)
    _patch_llm(monkeypatch)
    _FakeLLMClient.response = {
        "deck_title": "库存收尾",
        "page_label": "库存收尾",
        "pages": [
            {"layout": "cover", "title": "CLOSING & OUTLOOK", "subtitle": "两条主线收口"},
            {
                "layout": "cards",
                "title": "两条主线",
                "subtitle": "",
                "cards": [
                    {
                        "badge": "01",
                        "title": "合同转签至横琴旅文",
                        "points": ["对象：路秀新杰", "状态：已完成"],
                        "callout_title": "统一签约主体",
                        "callout_body": "后续结算与对账更清晰",
                    }
                ],
            },
            {"layout": "不存在的版式", "title": "兜底", "bullets": ["a", "b"]},
        ],
    }
    with _test_session() as db:
        _prepare(db)
        deck = generate_slides(
            _request(template_title="AI Lab 月度", page_label="库存收尾", page_count=3, skeleton_pages=["cover"]),
            db,
            _member(),
        )

    assert [page.layout for page in deck.pages] == ["cover", "cards", "bullets"]
    assert deck.pages[1].cards[0].callout_body == "后续结算与对账更清晰"
    # 传给模型的 system prompt / payload 里带上了调用方给的口径
    call = _FakeLLMClient.calls[0]
    # 总页数 3 - 勾选的封面 1 = 2 张内容页
    assert "内容页数量 = 2" in call["system"]
    assert call["payload"]["正文口述"].startswith("本月完成库存收尾")
    assert call["payload"]["模板标题"] == "AI Lab 月度"
    assert call["payload"]["总页数"] == "3"


def test_generate_fills_missing_skeleton_pages(monkeypatch: pytest.MonkeyPatch) -> None:
    """勾了骨架页但模型没给：用确定性的最小页面补上，不能悄悄少页。"""
    _patch_runtime_model(monkeypatch)
    _patch_llm(monkeypatch)
    _FakeLLMClient.response = {
        "deck_title": "年度汇报",
        "page_label": "",
        "pages": [{"layout": "bullets", "title": "正文一", "bullets": ["a"]}],
    }
    with _test_session() as db:
        _prepare(db)
        deck = generate_slides(
            _request(template_title="年度汇报", page_label="年度", skeleton_pages=["cover", "toc", "end"]),
            db,
            _member(),
        )

    assert [page.layout for page in deck.pages] == ["cover", "toc", "bullets", "closing"]
    assert deck.pages[0].title == "年度汇报"
    assert deck.pages[1].cards[0].title == "正文一"
    # 内容页的 label 为空时回落到请求里的页面文字
    assert deck.page_label == "年度"


def test_generate_raises_when_model_returns_no_pages(monkeypatch: pytest.MonkeyPatch) -> None:
    _patch_runtime_model(monkeypatch)
    _patch_llm(monkeypatch)
    _FakeLLMClient.response = {"pages": []}
    with _test_session() as db:
        _prepare(db)
        with pytest.raises(HTTPException) as excinfo:
            generate_slides(_request(), db, _member())
    assert excinfo.value.status_code == 502
    assert "没有返回任何页面" in str(excinfo.value.detail)


def test_generate_translates_llm_failure(monkeypatch: pytest.MonkeyPatch) -> None:
    _patch_runtime_model(monkeypatch)

    class _Boom:
        def __init__(self, model_config: Any) -> None:
            pass

        def generate_json(self, *_: Any, **__: Any) -> Any:
            raise agent_apps.LLMError("模型网关超时")

    monkeypatch.setattr(agent_apps, "LLMClient", _Boom)
    with _test_session() as db:
        _prepare(db)
        with pytest.raises(HTTPException) as excinfo:
            generate_slides(_request(), db, _member())
    assert excinfo.value.status_code == 502
    assert "模型调用失败" in str(excinfo.value.detail)


# ---------------------------------------------------------------- 选项对齐


def test_generate_drops_unchecked_skeleton_page(monkeypatch: pytest.MonkeyPatch) -> None:
    """没勾「首页」时模型擅自生成的封面页要丢掉，否则左侧选项形同虚设。"""
    _patch_runtime_model(monkeypatch)
    _patch_llm(monkeypatch)
    _FakeLLMClient.response = {
        "deck_title": "库存收尾",
        "pages": [
            {"layout": "cover", "title": "模型自作主张的封面"},
            {"layout": "bullets", "title": "正文一", "bullets": ["a"]},
            {"layout": "closing", "title": "模型自作主张的结尾"},
        ],
    }
    with _test_session() as db:
        _prepare(db)
        deck = generate_slides(_request(skeleton_pages=[]), db, _member())

    assert [page.layout for page in deck.pages] == ["bullets"]


def test_generate_truncates_content_pages_to_requested_count(monkeypatch: pytest.MonkeyPatch) -> None:
    """「页数」是总页数：勾了 2 张骨架页时，总 5 页 = 2 骨架 + 3 内容。"""
    _patch_runtime_model(monkeypatch)
    _patch_llm(monkeypatch)
    _FakeLLMClient.response = {
        "deck_title": "库存收尾",
        "pages": [
            {"layout": "bullets", "title": f"正文{i}", "bullets": ["a"]} for i in range(1, 6)
        ],
    }
    with _test_session() as db:
        _prepare(db)
        deck = generate_slides(_request(page_count=5, skeleton_pages=["cover", "end"]), db, _member())

    assert [page.layout for page in deck.pages] == ["cover", "bullets", "bullets", "bullets", "closing"]
    assert [page.title for page in deck.pages if page.layout == "bullets"] == ["正文1", "正文2", "正文3"]


def test_generate_keeps_one_content_page_when_skeleton_pages_fill_the_budget(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """总页数给得比骨架页还小：宁可总页数超出，也不能出一份没有内容页的 deck。"""
    _patch_runtime_model(monkeypatch)
    _patch_llm(monkeypatch)
    _FakeLLMClient.response = {
        "pages": [{"layout": "bullets", "title": "正文", "bullets": ["a"]}],
    }
    with _test_session() as db:
        _prepare(db)
        deck = generate_slides(_request(page_count=2, skeleton_pages=["cover", "toc", "end"]), db, _member())

    assert [page.layout for page in deck.pages] == ["cover", "toc", "bullets", "closing"]


def test_content_page_target_subtracts_skeleton_pages() -> None:
    assert slides_deck.content_page_target("custom", 5, ["cover", "end"]) == 3
    assert slides_deck.content_page_target("custom", 5, []) == 5
    assert slides_deck.content_page_target("custom", 1, ["cover"]) is not None
    assert slides_deck.content_page_target("custom", 1, ["cover"]) == 1
    assert slides_deck.content_page_target("auto", 5, ["cover"]) is None
    # 只认已知的骨架页 key，杂项不影响计算
    assert slides_deck.content_page_target("custom", 4, ["cover", "乱写"]) == 3


def test_generate_skips_truncation_in_auto_mode(monkeypatch: pytest.MonkeyPatch) -> None:
    _patch_runtime_model(monkeypatch)
    _patch_llm(monkeypatch)
    _FakeLLMClient.response = {
        "pages": [{"layout": "bullets", "title": f"正文{i}", "bullets": ["a"]} for i in range(1, 6)],
    }
    with _test_session() as db:
        _prepare(db)
        deck = generate_slides(
            _request(page_count_mode="auto", page_count=0, skeleton_pages=[]), db, _member()
        )

    assert len(deck.pages) == 5


# ---------------------------------------------------------------- 样式 prompt


def test_generate_appends_style_prompt_when_checked(monkeypatch: pytest.MonkeyPatch) -> None:
    """勾了「旅文汇报样式」：清单里的 prompt 会进 system prompt。"""
    _patch_runtime_model(monkeypatch)
    _patch_llm(monkeypatch)
    _FakeLLMClient.response = {"pages": [{"layout": "bullets", "title": "正文", "bullets": ["a"]}]}
    with _test_session() as db:
        _prepare(db)
        generate_slides(_request(style_prompt="结论先行，用词统一采用「复星旅文」口径。"), db, _member())

    system = _FakeLLMClient.calls[0]["system"]
    assert "复星旅文" in system
    assert "样式要求" in system


def test_generate_omits_style_section_without_style_prompt(monkeypatch: pytest.MonkeyPatch) -> None:
    _patch_runtime_model(monkeypatch)
    _patch_llm(monkeypatch)
    _FakeLLMClient.response = {"pages": [{"layout": "bullets", "title": "正文", "bullets": ["a"]}]}
    with _test_session() as db:
        _prepare(db)
        generate_slides(_request(), db, _member())

    assert "样式要求" not in _FakeLLMClient.calls[0]["system"]


def test_generate_tells_model_which_skeleton_pages_are_needed(monkeypatch: pytest.MonkeyPatch) -> None:
    _patch_runtime_model(monkeypatch)
    _patch_llm(monkeypatch)
    _FakeLLMClient.response = {"pages": [{"layout": "bullets", "title": "正文", "bullets": ["a"]}]}
    with _test_session() as db:
        _prepare(db)
        generate_slides(_request(skeleton_pages=["cover", "end"]), db, _member())

    system = _FakeLLMClient.calls[0]["system"]
    assert "封面页" in system and "结尾页" in system
    assert "目录页" not in system


# ---------------------------------------------------------------- 导出 pptx


def _deck_with_all_layouts():
    raw = {
        "deck_title": "AI 先锋项目",
        "page_label": "旅文汇报",
        "pages": [
            {"layout": "cover", "title": "AI 先锋项目", "subtitle": "四步路径成型"},
            {"layout": "toc", "title": "目录", "cards": [{"badge": "01", "title": "本月进展"}]},
            {"layout": "section", "title": "本月进展", "subtitle": "三条线同步推进"},
            {"layout": "metrics", "title": "关键指标", "cards": [{"badge": "397", "title": "大使人数"}]},
            {
                "layout": "cards",
                "title": "两条主线",
                "cards": [
                    {"badge": "01", "title": "合同转签", "points": ["对象：路秀新杰"]},
                    {"badge": "02", "title": "开票回收", "points": ["已完成：17 家"]},
                ],
            },
            {"layout": "bullets", "title": "问题与风险", "bullets": ["认证周期偏长"]},
            {"layout": "closing", "title": "谢谢"},
        ],
    }
    return slides_deck.finalize_deck(
        raw,
        template_title="AI 先锋项目",
        page_label="旅文汇报",
        skeleton_pages=["cover", "toc", "end"],
        page_count_mode="auto",
    )


def test_export_slides_returns_openable_pptx() -> None:
    deck = _deck_with_all_layouts()
    response = export_slides(SlidesExportRequest(deck=deck, file_name="AI 先锋项目汇报"))

    assert response.media_type == PPTX_MEDIA_TYPE
    assert bytes(response.body)[:2] == b"PK"  # zip 容器
    assert response.headers["content-length"] == str(len(response.body))

    # 能被 python-pptx 读回来，且页数一致
    presentation = Presentation(BytesIO(bytes(response.body)))
    assert len(presentation.slides) == len(deck.pages) == 7
    assert presentation.slide_width > presentation.slide_height  # 16:9


def test_export_slides_uses_rfc5987_file_name_for_chinese() -> None:
    response = export_slides(SlidesExportRequest(deck=_deck_with_all_layouts(), file_name="月度汇报"))
    disposition = response.headers["content-disposition"]
    assert "filename*=UTF-8''" in disposition
    assert "%E6%9C%88%E5%BA%A6%E6%B1%87%E6%8A%A5" in disposition  # 「月度汇报」
    # 纯中文名在 ASCII 兜底字段里回落成 slides.pptx，而不是被削成 ".pptx"
    assert disposition.startswith('attachment; filename="slides.pptx"')


def test_export_slides_falls_back_to_deck_title() -> None:
    deck = _deck_with_all_layouts()
    response = export_slides(SlidesExportRequest(deck=deck))
    assert "AI" in response.headers["x-file-name"]


def test_export_slides_rejects_empty_deck() -> None:
    with pytest.raises(HTTPException) as excinfo:
        export_slides(SlidesExportRequest(deck=slides_deck.SlidesDeck(deck_title="空的")))
    assert excinfo.value.status_code == 400
    assert "还没有可导出的幻灯片" in str(excinfo.value.detail)


def test_export_slides_sanitizes_unknown_layout() -> None:
    """直接打接口塞进来的越界版式也要被收敛，不能让 pptx 渲染器拿到未知 layout。"""
    deck = slides_deck.SlidesDeck(
        deck_title="兜底",
        pages=[slides_deck.SlidesPage(layout="随便写", title="正文", bullets=["a"])],
    )
    response = export_slides(SlidesExportRequest(deck=deck))
    presentation = Presentation(BytesIO(bytes(response.body)))
    assert len(presentation.slides) == 1


# ---------------------------------------------------------------- 归一化边界


def test_normalize_page_caps_lists_and_trims_text() -> None:
    page = slides_deck.normalize_page(
        {
            "layout": "CARDS ",
            "title": "  多   空格  ",
            "bullets": ["1", "", "2", "3", "4", "5", "6", "7"],
            "cards": [{"title": f"卡{index}"} for index in range(10)],
        }
    )
    assert page.layout == "cards"
    assert page.title == "多 空格"
    assert page.bullets == ["1", "2", "3", "4", "5", "6"]
    assert len(page.cards) == slides_deck.MAX_CARDS


def test_normalize_card_drops_empty_cards() -> None:
    page = slides_deck.normalize_page({"layout": "cards", "cards": [{}, {"title": "有效"}]})
    assert [card.title for card in page.cards] == ["有效"]


def test_normalize_card_caps_points_and_accepts_fields_alias() -> None:
    card = slides_deck.normalize_card({"fields": ["a", "b", "c", "d", "e"]})
    assert card.points == ["a", "b", "c", "d"]


def test_normalize_deck_truncates_long_text() -> None:
    deck = slides_deck.finalize_deck(
        {"pages": [{"layout": "bullets", "title": "字" * 200, "bullets": ["字" * 500]}]},
        template_title="",
        page_label="",
        skeleton_pages=[],
        page_count_mode="auto",
    )
    assert len(deck.pages[0].title) == slides_deck.MAX_TITLE_CHARS
    assert len(deck.pages[0].bullets[0]) == slides_deck.MAX_BODY_CHARS
    assert deck.pages[0].title.endswith("…")


def test_normalize_deck_falls_back_to_request_labels() -> None:
    deck = slides_deck.finalize_deck(
        {"pages": [{"layout": "bullets", "title": "正文"}]},
        template_title="AI Lab 月度",
        page_label="",
        skeleton_pages=[],
        page_count_mode="auto",
    )
    assert deck.deck_title == "AI Lab 月度"
    assert deck.page_label == "AI Lab 月度"
